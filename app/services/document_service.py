from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Mm
from app.database.models import Letter, Organization
def create_docx(letter:Letter, org:Organization|None, out:Path)->Path:
    d=Document(); sec=d.sections[0]; sec.top_margin=Mm(18); sec.bottom_margin=Mm(18); sec.left_margin=Mm(25); sec.right_margin=Mm(20)
    if org:
        if org.logo_path and Path(org.logo_path).exists(): d.add_picture(org.logo_path,width=Mm(28))
        p=d.add_paragraph(); p.alignment=1; r=p.add_run(org.full_name); r.bold=True
        details=" | ".join(x for x in [org.address,org.phone,org.email,org.website] if x)
        if details: p=d.add_paragraph(details); p.alignment=1
    d.add_paragraph(f"Sana: {date.today():%d.%m.%Y}    № {letter.letter_number or '—'}")
    p=d.add_paragraph(); p.alignment=2; p.add_run(letter.recipient).bold=True
    if letter.recipient_person: d.add_paragraph(letter.recipient_person).alignment=2
    p=d.add_paragraph(); p.add_run(f"Mavzu: {letter.subject}").bold=True
    for block in letter.content.split("\n"):
        if block.strip(): d.add_paragraph(block.strip())
    d.add_paragraph("")
    d.add_paragraph(f"{letter.signer_position or ''}                              {letter.signer_name or ''}")
    if letter.attachments: d.add_paragraph(f"Ilova: {letter.attachments}")
    out.parent.mkdir(parents=True,exist_ok=True); d.save(out); return out
